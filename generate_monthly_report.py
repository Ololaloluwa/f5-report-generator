#!/usr/bin/env python3
"""
MTN F5 MONTHLY Security Metrics Report - data + PowerPoint
============================================================

Builds the monthly deck from the same source workbook the weekly report
uses (MTN_F5_ATTACKS_SOURCE.xlsx) and the same PowerPoint template, so the
two reports look alike. Everything is worked out from the source every
time - including the previous months on the trend chart - so nobody types
last month's figures in by hand. Agreed with the user 28-Sep-2026:

  Month totals (Requests Inspected / WAF Blocks), per site:
    - a week that sits entirely inside the month  -> its "Weekly Request
      Throughput" / "Weekly WAF Attack Breakdown" tables
    - a week that crosses into the next/previous month -> only the days
      that belong to this month, from that sheet's "<Day> Request
      Throughput" / "<Day> WAF Attack Breakdown" tables
      (e.g. 28 Sep - 4 Oct gives Mon-Wed to September)
    - WAF Blocks = the attack-breakdown totals, so the PowerPoint and the
      Excel attack tables always agree
  Daily tables that don't add up to their weekly table are flagged
  ([WARN]) but the report is still produced.
  Trend chart: this month + up to 3 months before it (max 4), each only if
    the source covers every day of it.
  Pool member status: the last day of the month.
  Node changes: every day of the month compared with the day before.
  Certificates + the per-site pool table: as in the weekly report.

Nothing in the weekly scripts is changed by this file - it only imports
their helpers.
"""

import calendar
import copy
import math
import re
from collections import namedtuple
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

import openpyxl
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

import generate_weekly_report as wk
from generate_weekly_report import SITES, WEEKDAYS, apply_font, set_run_text

MAX_MONTHS = 4                 # months on the trend chart / Excel dashboard
KPI_JUMP_WARN_RATIO = 0.5      # warn if a KPI moved more than +-50% vs last month


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------

def num(v):
    """Numbers in the source are sometimes text ("1,234", "\xa01234") - read them all as int."""
    if v is None or isinstance(v, bool):
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).replace("\xa0", "").replace(",", "").strip()
    if not s or s.startswith("#"):
        return 0
    try:
        return int(float(s))
    except ValueError:
        return 0


def month_days(ym):
    y, m = ym
    return [date(y, m, d) for d in range(1, calendar.monthrange(y, m)[1] + 1)]


def month_name(ym, short=False):
    return (calendar.month_abbr if short else calendar.month_name)[ym[1]]


def month_label(ym):
    return f"{month_name(ym)} {ym[0]}"


def prev_month(ym):
    y, m = ym
    return (y - 1, 12) if m == 1 else (y, m - 1)


def parse_month(text):
    """'2026-09' / '2026-9' / 'Sep 2026' / 'September 2026' -> (2026, 9)."""
    text = str(text).strip()
    m = re.match(r"^(\d{4})-(\d{1,2})$", text)
    if m:
        return int(m.group(1)), int(m.group(2))
    m = re.match(r"^([A-Za-z]+)\.?\s+(\d{4})$", text)
    if m:
        mon = wk._resolve_month(m.group(1))
        if mon:
            return int(m.group(2)), mon
    raise ValueError(f"Can't read the month '{text}' - use e.g. 2026-09 or 'Sep 2026'")


def _short(d):
    return f"{d.day} {d:%b}"


def _day_name(d):
    return WEEKDAYS[d.weekday()]


def _norm_attack(name):
    s = str(name).lower().replace("-", " ").replace("/", " ")
    return re.sub(r"\s+", " ", s).strip()


def _get(table, key):
    """Case/space-insensitive lookup in a read_kv_table() result."""
    want = key.lower()
    for k, v in table.items():
        if str(k).strip().lower() == want:
            return v
    return None


def _pct_change(cur, old):
    if not old:
        return None
    return (cur - old) / old * 100


def _arrow(pct):
    if pct is None:
        return "n/a"
    sign = "▲" if pct >= 0 else "▼"
    return f"{sign} {abs(pct):.1f}%"


# --------------------------------------------------------------------------
# A light, read-only stand-in for an openpyxl worksheet
# --------------------------------------------------------------------------
# A month needs up to ~19 weeks x 3 sites of sheets. Copying those into an
# openpyxl workbook (what the weekly report does for its 12 sheets) would
# need several hundred MB on the free Render plan, so each sheet is kept as
# a plain list of row tuples instead. It answers the same calls the weekly
# table readers make (ws.cell(row, column).value, ws.max_row, iter_rows),
# so those readers are reused as-is.

_Cell = namedtuple("_Cell", "value")


class Grid:
    def __init__(self, title, rows):
        self.title = title
        self._rows = rows

    @property
    def max_row(self):
        return len(self._rows)

    def cell(self, row, column):
        try:
            return _Cell(self._rows[row - 1][column - 1])
        except IndexError:
            return _Cell(None)

    def iter_rows(self, values_only=True, **_kw):
        return iter(self._rows)


# --------------------------------------------------------------------------
# Which months the source covers
# --------------------------------------------------------------------------

@dataclass
class MonthCoverage:
    ym: tuple
    missing: dict               # {site: [dates with no sheet]}

    @property
    def complete(self):
        return not any(self.missing.values())

    @property
    def any_data(self):
        n = len(month_days(self.ym))
        return any(len(v) < n for v in self.missing.values())


def month_coverage(weeks, ym):
    missing = {}
    for site in SITES:
        covered = set()
        for (start, end), refs in weeks.items():
            if site in refs:
                covered.update(start + timedelta(days=i) for i in range(7))
        missing[site] = [d for d in month_days(ym) if d not in covered]
    return MonthCoverage(ym, missing)


def available_months(weeks, today=None):
    """Every month any (already started) week touches, oldest first, with its coverage."""
    today = today or date.today()
    yms = set()
    for start, end in weeks:
        if start > today:
            continue            # a future-dated typo mustn't add a month
        for i in range(7):
            d = start + timedelta(days=i)
            yms.add((d.year, d.month))
    return [month_coverage(weeks, ym) for ym in sorted(yms)]


def pick_target_month(covs, month=None):
    """The requested month, else the newest month the source covers completely
    (else the newest month with any data at all)."""
    if not covs:
        raise ValueError("No parsable weekly sheets found in workbook.")
    if month is not None:
        for c in covs:
            if c.ym == tuple(month):
                return c
        have = ", ".join(month_label(c.ym) for c in covs)
        raise ValueError(f"The source has no sheets for {month_label(month)} (it has: {have})")
    complete = [c for c in covs if c.complete]
    return (complete or covs)[-1]


def calendar_weeks(ym):
    """Every Monday-Sunday week that has at least one day in the month, whether or not the source has it."""
    first, last = month_days(ym)[0], month_days(ym)[-1]
    monday = wk.monday_of(first)
    out = []
    while monday <= last:
        out.append((monday, monday + timedelta(days=6)))
        monday += timedelta(days=7)
    return out


def weeks_for_month(weeks, ym):
    first, last = month_days(ym)[0], month_days(ym)[-1]
    return sorted(k for k in weeks if k[0] <= last and k[1] >= first)


# --------------------------------------------------------------------------
# Loading the source
# --------------------------------------------------------------------------

@dataclass
class MonthContext:
    target: tuple                   # (year, month) being reported
    coverage: MonthCoverage
    trend: list                     # months on the chart, oldest -> newest (incl. target)
    skipped: list                   # (ym, reason) months left off the chart
    weeks: dict                     # all readable weeks {key: {site: SheetRef}}
    grids: dict                     # {sheet_name: Grid} - only the sheets we need
    source_sheetnames: list
    latest_week: tuple


def load_month_context(path, month=None, n_months=MAX_MONTHS):
    """
    Streams only the sheets the chosen months need, values only. Weeks that
    touch the reported month are kept whole (the node lists sit far to the
    right); older weeks only keep columns A-L, which is where every table
    the report reads lives.
    """
    src = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        weeks = wk.group_sheets_by_week(src)
        covs = available_months(weeks)
        target_cov = pick_target_month(covs, month)
        target = target_cov.ym

        by_ym = {c.ym: c for c in covs}
        trend, skipped = [target], []
        ym = target
        for _ in range(n_months - 1):
            ym = prev_month(ym)
            cov = by_ym.get(ym) or month_coverage(weeks, ym)
            if cov.complete:
                trend.insert(0, ym)
            elif cov.any_data:
                skipped.append((ym, "the source doesn't cover every day of it"))
        full_keys = set(weeks_for_month(weeks, target))
        need_keys = set()
        for ym in trend:
            need_keys.update(weeks_for_month(weeks, ym))

        grids = {}
        for key in sorted(need_keys):
            for ref in weeks[key].values():
                ws = src[ref.sheet_name]
                ws.reset_dimensions()          # don't trust the stored sheet size
                if key in full_keys:
                    rows = list(ws.iter_rows(values_only=True))
                else:
                    rows = list(ws.iter_rows(max_col=12, values_only=True))
                grids[ref.sheet_name] = Grid(ref.sheet_name, rows)
        return MonthContext(
            target=target, coverage=target_cov, trend=trend, skipped=skipped, weeks=weeks,
            grids=grids, source_sheetnames=list(src.sheetnames),
            latest_week=wk.latest_real_week(weeks),
        )
    finally:
        src.close()


# --------------------------------------------------------------------------
# Table readers (daily + weekly)
# --------------------------------------------------------------------------

def _safe(value, fallback):
    if isinstance(value, str) and value.strip().startswith("#"):
        return fallback
    return num(value) if value is not None else fallback


def read_attack_rows(ws, title):
    """
    Rows of a "<...> WAF Attack Breakdown" table: attack type, violation
    count, blocked, allowed, remediated, IPs. The header row is searched
    for BELOW the title (the weekly reader searches from the top of the
    sheet, which only works for the weekly table at the top). Returns None
    if the title isn't in the sheet at all, [] if the table is empty.
    """
    title_row = wk.find_summary_row(ws, title)
    if title_row is None:
        return None
    header = None
    for r in range(title_row, title_row + 10):
        v = ws.cell(r, 2).value
        if isinstance(v, str) and v.strip().lower() == "attack type":
            header = r
            break
    if header is None:
        return []
    rows = []
    r = header + 1
    while r < header + 200:
        name = ws.cell(r, 2).value
        if name is None or not str(name).strip():
            break
        if str(name).strip().upper() != "TOTAL":
            count = num(ws.cell(r, 3).value)
            rows.append({
                "attack_type": str(name).strip(),
                "violation_count": count,
                "blocked": _safe(ws.cell(r, 4).value, count),
                "allowed": _safe(ws.cell(r, 5).value, 0),
                "remediated": _safe(ws.cell(r, 6).value, count),
                "ips": ws.cell(r, 7).value,
            })
        r += 1
    return rows


def read_requests(ws, title):
    """Requests Inspected from a "<...> Request Throughput" table, or None if the table isn't there."""
    table = wk.read_kv_table(ws, title)
    if not table:
        return None
    return num(_get(table, "Requests Inspected"))


POOL_KEYS = ("Available", "Unavailable", "Offline", "Unknown", "TOTAL")


def _pool_from(table):
    out = {k.lower(): num(_get(table, k)) for k in POOL_KEYS}
    if not out["total"]:
        out["total"] = out["available"] + out["unavailable"] + out["offline"] + out["unknown"]
    return out


def read_pool(ws, day):
    """
    The "<Day> Pool Member Status" counts, or None. Thursday's title is
    sometimes missing in the source, so if the title isn't found and the
    sheet has exactly seven pool tables, the day's table is picked by
    position (Monday = 1st ... Sunday = 7th).
    """
    table = wk.read_kv_table(ws, f"{day} Pool Member Status")
    if table:
        return _pool_from(table)
    headers = []
    for r in range(1, ws.max_row + 1):
        for c in range(1, 10):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().lower() == "pool member status":
                nxt = ws.cell(r, c + 1).value
                if isinstance(nxt, str) and nxt.strip().lower() == "count":
                    headers.append((r, c))
                break
    if len(headers) != 7:
        return None
    r, c = headers[WEEKDAYS.index(day)]
    table = {}
    r += 1
    while True:
        k = ws.cell(r, c).value
        if k is None:
            break
        table[str(k).strip()] = ws.cell(r, c + 1).value
        r += 1
    return _pool_from(table) if table else None


# --------------------------------------------------------------------------
# Month metrics
# --------------------------------------------------------------------------

@dataclass
class WeekPart:
    key: tuple                 # (Monday, Sunday) of the sheet's week
    days: list                 # the dates of that week inside the month
    requests: int = 0
    waf_blocks: int = 0
    source: str = "weekly"     # "weekly" tables or "daily" tables
    missing: bool = False      # no sheet for this site/week

    @property
    def label(self):
        s, e = self.key
        text = f"{_short(s)} – {_short(e)}"
        if len(self.days) < 7:
            a, b = self.days[0], self.days[-1]
            text += f" ({a.day}–{b.day} {b:%b})" if a != b else f" ({_short(a)})"
        return text


@dataclass
class MonthMetrics:
    ym: tuple
    site: str = None                     # None = all sites combined
    requests_inspected: int = 0
    waf_blocks: int = 0
    pool_available: int = 0
    pool_unavailable: int = 0
    pool_offline: int = 0
    pool_unknown: int = 0
    pool_total: int = 0
    pool_date: date = None
    attack_rows: dict = field(default_factory=dict)   # {normalized name: row dict}
    parts: list = field(default_factory=list)         # WeekPart per week (per site only)
    daily: dict = field(default_factory=dict)         # {date: (requests, waf_blocks)}

    @property
    def label(self):
        return month_label(self.ym)

    @property
    def attack_counts(self):
        return {r["attack_type"]: r["violation_count"] for r in self.attack_rows.values()}

    @property
    def block_rate(self):
        return self.waf_blocks / self.requests_inspected * 100 if self.requests_inspected else 0

    @property
    def pool_pct(self):
        return self.pool_available / self.pool_total * 100 if self.pool_total else 0

    @property
    def peak_day(self):
        if not self.daily:
            return None
        d = max(self.daily, key=lambda k: self.daily[k][1])
        return d, self.daily[d][1]


def _merge_attack_rows(into, rows):
    for row in rows:
        key = _norm_attack(row["attack_type"])
        cur = into.get(key)
        if cur is None:
            cur = into[key] = {"attack_type": row["attack_type"], "violation_count": 0,
                               "blocked": 0, "allowed": 0, "remediated": 0, "ips": []}
        for f in ("violation_count", "blocked", "allowed", "remediated"):
            cur[f] += row[f]
        seen = set(cur["ips"])
        for ip in _ip_entries(row.get("ips")):
            if ip not in seen:
                seen.add(ip)
                cur["ips"].append(ip)


def _ip_entries(text):
    """
    "1.2.3.4, 5.6.7.8 (1,234), 9.9.9.9 (x3)" -> ["1.2.3.4", "5.6.7.8", "9.9.9.9"].
    Some tables put a hit count in brackets after an IP; that count is for one
    week/day only, so it's dropped from the month's combined list (and its
    comma doesn't split anything).
    """
    out = []
    for e in re.findall(r"[^,;()\n]+(?:\([^)]*\))?", str(text or "")):
        e = re.sub(r"\(.*?\)", "", e).strip().strip("'").strip()
        if e and e.upper() not in ("N/A", "NONE"):
            out.append(e)
    return out


def compute_site_month(ctx, site, ym, detail=False, warn=None):
    """
    One site's month. With detail=True (the month being reported) it also
    reads every daily table for the daily-vs-weekly check, the peak day,
    and the last-day pool status. `warn(text)` receives [WARN] texts; for
    an earlier month (detail=False) passing `warn` checks just the split
    weeks, since those are the only ones whose daily tables get used.
    """
    warn_given = warn is not None
    warn = warn or (lambda _t: None)
    mm = MonthMetrics(ym=ym, site=site)
    days_in_month = set(month_days(ym))
    for key in calendar_weeks(ym):
        start, end = key
        week_days = [start + timedelta(days=i) for i in range(7)]
        in_month = [d for d in week_days if d in days_in_month]
        ref = ctx.weeks.get(key, {}).get(site)
        part = WeekPart(key=key, days=in_month)
        mm.parts.append(part)
        if ref is None:
            part.missing = True
            continue
        ws = ctx.grids[ref.sheet_name]
        wlabel = f"{_short(start)} – {_short(end)}"

        daily = {}
        split = len(in_month) < 7
        if detail or split:
            # all 7 days whenever the daily/weekly cross-check will run
            for d in (week_days if (detail or warn_given) else in_month):
                dn = _day_name(d)
                req = read_requests(ws, f"{dn} Request Throughput")
                rows = read_attack_rows(ws, f"{dn} WAF Attack Breakdown")
                daily[d] = (req, rows)

        if len(in_month) == 7:
            part.source = "weekly"
            part.requests = num(read_requests(ws, "Weekly Request Throughput"))
            rows = read_attack_rows(ws, "Weekly WAF Attack Breakdown") or []
            part.waf_blocks = sum(r["violation_count"] for r in rows)
            _merge_attack_rows(mm.attack_rows, rows)
        else:
            part.source = "daily"
            for d in in_month:
                req, rows = daily[d]
                if req is None:
                    warn(f"{site} {wlabel}: no '{_day_name(d)} Request Throughput' table - "
                         f"{_short(d)}'s requests counted as 0")
                if rows is None:
                    warn(f"{site} {wlabel}: no '{_day_name(d)} WAF Attack Breakdown' table - "
                         f"{_short(d)}'s WAF blocks counted as 0")
                part.requests += num(req)
                part.waf_blocks += sum(r["violation_count"] for r in rows or [])
                _merge_attack_rows(mm.attack_rows, rows or [])
        mm.requests_inspected += part.requests
        mm.waf_blocks += part.waf_blocks

        if detail:
            for d in in_month:
                req, rows = daily[d]
                mm.daily[d] = (num(req), sum(r["violation_count"] for r in rows or []))
            _check_daily_vs_weekly(site, wlabel, ws, daily, len(in_month) == 7, warn)
        elif split and warn_given:
            # an earlier month on the trend chart: only its split weeks use daily tables
            _check_daily_vs_weekly(site, wlabel, ws, daily, False, warn)

    if detail:
        _read_last_day_pool(ctx, site, ym, mm, warn)
    return mm


def _check_daily_vs_weekly(site, wlabel, ws, daily, whole_week, warn):
    """Flags a week whose 7 daily tables don't add up to its weekly tables - one line per week (the report still runs)."""
    if len(daily) != 7 or any(req is None or rows is None for req, rows in daily.values()):
        return
    used = "weekly tables used" if whole_week else "daily tables used for this month's days"
    diffs = []
    w_req = read_requests(ws, "Weekly Request Throughput")
    d_req = sum(num(req) for req, _ in daily.values())
    if w_req and d_req != w_req:
        diffs.append(f"requests: days add up to {d_req:,}, weekly table says {w_req:,}")
    w_rows = read_attack_rows(ws, "Weekly WAF Attack Breakdown")
    if w_rows:
        w_blk = sum(r["violation_count"] for r in w_rows)
        d_blk = sum(sum(r["violation_count"] for r in rows) for _, rows in daily.values())
        if d_blk != w_blk:
            diffs.append(f"WAF blocks: days add up to {d_blk:,}, weekly table says {w_blk:,}")
    if diffs:
        warn(f"{site} {wlabel}: daily and weekly tables disagree - " + "; ".join(diffs) + f" ({used})")


def _read_last_day_pool(ctx, site, ym, mm, warn):
    """Pool status on the last day of the month; falls back to the latest earlier day that has it."""
    last = month_days(ym)[-1]
    d = last
    while d.month == ym[1]:
        key = (wk.monday_of(d), wk.monday_of(d) + timedelta(days=6))
        ref = ctx.weeks.get(key, {}).get(site)
        pool = None
        if ref is not None and ref.sheet_name in ctx.grids:
            pool = read_pool(ctx.grids[ref.sheet_name], _day_name(d))
        if pool and pool["total"]:
            mm.pool_available, mm.pool_unavailable = pool["available"], pool["unavailable"]
            mm.pool_offline, mm.pool_unknown, mm.pool_total = pool["offline"], pool["unknown"], pool["total"]
            mm.pool_date = d
            if d != last:
                warn(f"{site}: no Pool Member Status for {_day_name(last)} {_short(last)} (last day of "
                     f"the month) - used {_day_name(d)} {_short(d)} instead")
            return
        d -= timedelta(days=1)
    warn(f"{site}: no Pool Member Status table found for any day of {month_label(ym)} - pool shown as 0")


def combine(ym, site_months):
    total = MonthMetrics(ym=ym)
    for sm in site_months:
        total.requests_inspected += sm.requests_inspected
        total.waf_blocks += sm.waf_blocks
        for f in ("pool_available", "pool_unavailable", "pool_offline", "pool_unknown", "pool_total"):
            setattr(total, f, getattr(total, f) + getattr(sm, f))
        _merge_attack_rows(total.attack_rows, [dict(r, ips=", ".join(r["ips"])) for r in sm.attack_rows.values()])
        for d, (req, blk) in sm.daily.items():
            a, b = total.daily.get(d, (0, 0))
            total.daily[d] = (a + req, b + blk)
    dates = [sm.pool_date for sm in site_months if sm.pool_date]
    total.pool_date = max(dates) if dates else None
    return total


# --------------------------------------------------------------------------
# Node changes across the whole month
# --------------------------------------------------------------------------

def month_node_changes(ctx, site, ym):
    """
    [(date, name, ip, old, new)] for the month, plus ambiguous rows. Same
    rules as the weekly report (a node only counts if it's in the list on
    both days and its status changed), but every day of the month is
    compared with the day before - including Sunday -> Monday across two
    weekly sheets.
    """
    states = {}
    for key in weeks_for_month(ctx.weeks, ym):
        ref = ctx.weeks[key].get(site)
        if ref is None:
            continue
        per_day = wk.read_daily_node_states(ctx.grids[ref.sheet_name])
        for day, st in per_day.items():
            states[key[0] + timedelta(days=wk.DAYS.index(day))] = st
    changes, ambiguous = [], []
    for d in sorted(states):
        if d.month != ym[1] or d.year != ym[0]:
            continue
        prev = states.get(d - timedelta(days=1))
        if prev is None:
            continue
        cur = states[d]
        for node in sorted(set(prev) & set(cur)):
            before, after = prev[node], cur[node]
            if before == after:
                continue
            if len(before) > 1 or len(after) > 1:
                ambiguous.append((d, node))
                continue
            name, ip, _port = node
            changes.append((d, name, ip, next(iter(before)), next(iter(after))))
    return changes, ambiguous


# --------------------------------------------------------------------------
# PowerPoint
# --------------------------------------------------------------------------

GREY = RGBColor(0x59, 0x59, 0x59)


def _kpi_value_size(text):
    n = len(text)
    return Pt(26) if n <= 9 else Pt(23) if n <= 10 else Pt(20)


def _add_caption(slide, value_shape, label_shape, text):
    """Small grey line under a KPI card's label (e.g. '▲ 6.8% vs Jul')."""
    top = label_shape.top + label_shape.height - Emu(int(Inches(0.03)))
    box = slide.shapes.add_textbox(value_shape.left, top, value_shape.width, Emu(int(Inches(0.2))))
    tf = box.text_frame
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = label_shape.text_frame.paragraphs[0].alignment
    r = p.add_run()
    r.text = text
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    apply_font(tf)


def fill_month_footer(slide, ym):
    for shape in slide.shapes:
        if shape.has_text_frame and "CONFIDENTIAL" in shape.text_frame.text:
            runs = shape.text_frame.paragraphs[0].runs
            if runs:
                joined = "".join(r.text for r in runs)
                runs[0].text = re.sub(r"\[.*?\]", f"[{month_label(ym).upper()}]", joined)
                for extra in runs[1:]:
                    extra.text = ""
                apply_font(shape.text_frame)


def _set_subtitle(slide, text):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.startswith("Attack category"):
            set_run_text(shape, text)


def fill_month_dashboard(slide, cur, trend, prev):
    """
    The 4 KPI cards + charts for one site (or TOTAL). trend = [MonthMetrics]
    oldest -> newest (up to 4 months); prev = last month's MonthMetrics or None.
    """
    pool_txt = f"{cur.pool_pct:.2f}%"
    kpis = {
        "Requests Inspected": f"{cur.requests_inspected:,}",
        "WAF Blocks": f"{cur.waf_blocks:,}",
        "Pool Member Avail.": pool_txt,
        "F5 HA Uptime": "100%",
    }
    pm = month_name(prev.ym, short=True) if prev else None
    captions = {
        "Requests Inspected": f"{_arrow(_pct_change(cur.requests_inspected, prev.requests_inspected))} vs {pm}"
        if prev else "first month on record",
        "WAF Blocks": f"{_arrow(_pct_change(cur.waf_blocks, prev.waf_blocks))} vs {pm}  ·  "
                      f"{cur.block_rate:.1f}% of requests" if prev else f"{cur.block_rate:.1f}% of requests",
        "Pool Member Avail.": f"on {_short(cur.pool_date)}" if cur.pool_date else "",
        "F5 HA Uptime": "whole month",
    }

    shapes = list(slide.shapes)
    fp_shape = dough_shape = None
    for shape in shapes:
        if shape.has_chart and shape.chart.has_title:
            title = shape.chart.chart_title.text_frame.text
            if "False Positive" in title:
                fp_shape = shape
            elif "Pool Member Status" in title:
                dough_shape = shape
    if fp_shape is not None and dough_shape is not None:
        dough_shape.width = (fp_shape.left + fp_shape.width) - dough_shape.left
        fp_shape._element.getparent().remove(fp_shape._element)
        shapes = list(slide.shapes)

    for i, shape in enumerate(shapes):
        if shape.has_text_frame:
            label = shape.text_frame.text.strip()
            if label in kpis:
                value_shape = shapes[i - 1]
                set_run_text(value_shape, kpis[label])
                apply_font(value_shape.text_frame, size=_kpi_value_size(kpis[label]))
                if captions[label]:
                    _add_caption(slide, value_shape, shape, captions[label])
        if not shape.has_chart:
            continue
        chart = shape.chart
        title = chart.chart_title.text_frame.text if chart.has_title else ""
        if "Attack Category" in title:
            top10 = sorted(cur.attack_counts.items(), key=lambda kv: kv[1], reverse=True)[:10]
            top10 = list(reversed(top10)) or [("No data for this month", 0)]
            cd = CategoryChartData()
            cd.categories = [k for k, _ in top10]
            cd.add_series("Blocks", [v for _, v in top10])
            chart.replace_data(cd)
        elif "Request Rate vs Block Rate" in title:
            tf = chart.chart_title.text_frame
            runs = [r for p in tf.paragraphs for r in p.runs]
            if runs:
                runs[0].text = "Requests vs WAF Blocks (Monthly)"
                for r in runs[1:]:
                    r.text = ""
            cd = CategoryChartData()
            cd.categories = [month_name(t.ym, short=True).upper() for t in trend]
            cd.add_series("Requests Inspected", [t.requests_inspected for t in trend])
            cd.add_series("WAF Blocks", [t.waf_blocks for t in trend])
            chart.replace_data(cd)
            try:
                axis = chart.value_axis
                axis.tick_labels.number_format = '#,##0.0,,"M"'
                axis.tick_labels.number_format_is_linked = False
            except (ValueError, AttributeError):
                pass
        elif "Pool Member Status" in title:
            total = cur.pool_total or 1
            vals = [cur.pool_available, cur.pool_unavailable, cur.pool_offline, cur.pool_unknown]
            pcts = [v / total * 100 for v in vals]
            cd = CategoryChartData()
            cd.categories = [f"{n} ({p:.0f}%)" for n, p in zip(("Available", "Unavailable", "Offline", "Unknown"), pcts)]
            cd.add_series("Pool Health", pcts)
            chart.replace_data(cd)
            wk.hide_all_data_labels(chart)
            if cur.pool_date:
                tf = chart.chart_title.text_frame
                runs = [r for p in tf.paragraphs for r in p.runs]
                if runs:
                    runs[0].text = f"Pool Member Status ({_short(cur.pool_date)})"
                    for r in runs[1:]:
                        r.text = ""


def _table(slide, rows, left, top, width, ratios, size=Pt(10), row_h=Inches(0.23), bold_last=True):
    """A plain table in the template's default table style; header row + optional bold last row."""
    n_rows, n_cols = len(rows), len(rows[0])
    gf = slide.shapes.add_table(n_rows, n_cols, Emu(int(left)), Emu(int(top)),
                                Emu(int(width)), Emu(int(row_h * n_rows)))
    table = gf.table
    s = sum(ratios)
    for c, ratio in enumerate(ratios):
        table.columns[c].width = Emu(int(width * ratio / s))
    for r in range(n_rows):
        table.rows[r].height = Emu(int(row_h))
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.text = str(rows[r][c]) or " "      # a blank cell still needs a run to take the font size
            cell.margin_top = cell.margin_bottom = Emu(int(Inches(0.02)))
            cell.margin_left = cell.margin_right = Emu(int(Inches(0.06)))
            apply_font(cell.text_frame, size=size, bold=(r == 0 or (bold_last and r == n_rows - 1)))
    return gf


def _pool_rows(cur, breakdown=None):
    sites = [s for s in SITES if breakdown and s in breakdown]
    head = (["Status"] if sites else ["Pool Member Status"]) + sites + (["Total"] if sites else ["Count"]) + ["%"]
    total = cur.pool_total or 1
    rows = [head]
    for label, attr in (("Available", "pool_available"), ("Unavailable", "pool_unavailable"),
                        ("Offline", "pool_offline"), ("Unknown", "pool_unknown")):
        v = getattr(cur, attr)
        rows.append([label] + [f"{getattr(breakdown[s], attr):,}" for s in sites]
                    + [f"{v:,}", f"{v / total * 100:.2f}%"])
    rows.append(["TOTAL"] + [f"{breakdown[s].pool_total:,}" for s in sites] + [f"{cur.pool_total:,}", "100%"])
    return rows


def _summary_lines(cur, prev, node_text, site_breakdown=None):
    lines = [("", "We had 100% Up time on the active and standby nodes."),
             ("Incidents:", " None recorded")]
    lines.append(("Requests Inspected:", f" {cur.requests_inspected:,}"))
    lines.append(("WAF Blocks:", f" {cur.waf_blocks:,}"))
    lines.append(("Block rate:", f" {cur.block_rate:.1f}% of requests"))
    if prev:
        lines.append((f"vs {month_name(prev.ym)}:",
                      f" requests {_arrow(_pct_change(cur.requests_inspected, prev.requests_inspected))}"
                      f" · WAF blocks {_arrow(_pct_change(cur.waf_blocks, prev.waf_blocks))}"))
    peak = cur.peak_day
    if peak and peak[1]:
        d, blk = peak
        lines.append(("Peak attack day:", f" {d:%a} {_short(d)} – {blk:,} WAF blocks"))
    when = f" ({_short(cur.pool_date)})" if cur.pool_date else ""
    lines.append((f"Pool Availability{when}:", f" {cur.pool_pct:.2f}%"))
    if node_text:
        lines.append(("Node changes:", f" {node_text}"))
    return lines


def fill_month_exec_slide(prs, slide, cur, prev, node_text, site_breakdown=None):
    """
    Left: the month in words (incidents line is for the engineer to edit).
    Right: pool status on the last day (+ per site on the TOTAL slide), and
    below it the month week by week (per site) or site by site (TOTAL).
    """
    footer_top = wk._footer_top(prs, slide)
    top = int(prs.slide_height * 0.21)
    left = int(Inches(0.4))
    text_w = int(Inches(4.55))
    lines = _summary_lines(cur, prev, node_text)
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(text_w), Emu(footer_top - top - int(Inches(0.1))))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (label, value) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(3)
        if label:
            r1 = p.add_run()
            r1.text = label
            r1.font.bold = True
        r2 = p.add_run()
        r2.text = value
    apply_font(tf, size=Pt(12))

    right = int(Inches(5.15))
    right_w = int(prs.slide_width - right - Inches(0.4))
    pool = _pool_rows(cur, site_breakdown)
    if site_breakdown:
        ratios = [0.27] + [0.14] * 3 + [0.15, 0.16]
    else:
        ratios = [0.46, 0.26, 0.28]
    gf = _table(slide, pool, right, top, right_w, ratios)
    y = top + gf.height + int(Inches(0.2))

    if site_breakdown:
        rows = [["Site", "Requests", "WAF Blocks", "Blocked %"]]
        for s in SITES:
            m = site_breakdown[s]
            rows.append([s, f"{m.requests_inspected:,}", f"{m.waf_blocks:,}", f"{m.block_rate:.1f}%"])
        rows.append(["TOTAL", f"{cur.requests_inspected:,}", f"{cur.waf_blocks:,}", f"{cur.block_rate:.1f}%"])
        _table(slide, rows, right, y, right_w, [0.17, 0.3, 0.3, 0.23])
    else:
        rows = [["Week (days in month)", "Requests", "WAF Blocks"]]
        for part in cur.parts:
            if part.missing:
                rows.append([part.label, "no sheet", "no sheet"])
            else:
                rows.append([part.label, f"{part.requests:,}", f"{part.waf_blocks:,}"])
        rows.append(["TOTAL", f"{cur.requests_inspected:,}", f"{cur.waf_blocks:,}"])
        size = Pt(10) if len(rows) <= 8 else Pt(9)
        _table(slide, rows, right, y, right_w, [0.46, 0.27, 0.27], size=size,
               row_h=min(Inches(0.23), (footer_top - y - Inches(0.08)) / len(rows)))


NODE_ROWS_PER_SLIDE = 13


def build_month_node_slides(prs, template_exec_slide, site, ym, changes):
    """
    The month's node changes as a table (date / node / change), 13 rows a
    slide - a month has too many for the weekly report's free-text list.
    The date is only written on the first row of each day.
    """
    rows, last = [], None
    for d, name, ip, old, new in changes:
        change = f"{old.title()} → {new.title()}" + (" (recovered)" if new == "AVAILABLE" else "")
        rows.append([f"{d:%a} {d.day} {d:%b}" if d != last else "", f"{name} ({ip})", change])
        last = d
    pages = [rows[i:i + NODE_ROWS_PER_SLIDE] for i in range(0, len(rows), NODE_ROWS_PER_SLIDE)]
    slides = []
    for p, page in enumerate(pages, start=1):
        if page[0][0] == "":                     # a day carried over from the previous slide
            i = (p - 1) * NODE_ROWS_PER_SLIDE
            while not rows[i][0]:
                i -= 1
            page[0] = [rows[i][0] + " (cont.)"] + page[0][1:]
        suffix = f" ({p}/{len(pages)})" if len(pages) > 1 else ""
        slide = wk.clone_header_slide(
            prs, template_exec_slide, f"F5 BIG-IP — {site} NODE CHANGES{suffix}",
            f"Pool member status changes during {month_label(ym)} (each day compared with the day before)"
            f" · {len(changes)} in total")
        left, top, width, height = wk._content_box(prs, slide)
        left, width = int(Inches(0.45)), int(prs.slide_width - 2 * Inches(0.45))
        row_h = min(Inches(0.26), height / (NODE_ROWS_PER_SLIDE + 1))
        _table(slide, [["Date", "Node (IP)", "Change"]] + page, left, top, width, [0.14, 0.56, 0.30],
               size=Pt(10), row_h=row_h, bold_last=False)
        slides.append(slide)
    return slides


def fill_month_cover(slide, ym):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        paras = shape.text_frame.paragraphs
        text = shape.text_frame.text
        if text.startswith("F5 BIG-IP") and len(paras) >= 2:
            for para, new in ((paras[0], f"F5 BIG-IP {month_name(ym).upper()}"),
                              (paras[1], f"{ym[0]} Monthly Metrics")):
                if para.runs:
                    para.runs[0].text = new
                    for extra in para.runs[1:]:
                        extra.text = ""
            apply_font(shape.text_frame)
            return


def fill_monthly_pptx(template_path, out_path, ym, total_trend, site_trends, node_changes=None,
                      certs=None):
    """
    total_trend: [MonthMetrics] oldest -> newest, all sites combined (last = reported month).
    site_trends: {site: [MonthMetrics]} the same months per site.
    """
    prs = Presentation(template_path)
    slides = list(prs.slides)
    template_exec = slides[2]          # untouched copy for cloning extra slides
    extra = {s: [] for s in SITES}
    for site in SITES:
        if node_changes and node_changes.get(site):
            extra[site] += build_month_node_slides(prs, template_exec, site, ym, node_changes[site])
        if certs and site in certs:
            for kind in ("ssl", "device"):
                extra[site] += wk.build_cert_slides(prs, template_exec, site, kind,
                                                    certs[site][kind], certs[site], ym[0])

    def prev_of(trend):
        return trend[-2] if len(trend) >= 2 and trend[-2].ym == prev_month(ym) else None

    fill_month_cover(slides[0], ym)
    subtitle = f"Monthly summary · {month_label(ym)}"
    positions = [(1, 2, "IKY"), (3, 4, "VGC"), (5, 6, "OJT")]
    for dash, ex, site in positions:
        trend = site_trends[site]
        cur = trend[-1]
        fill_month_dashboard(slides[dash], cur, trend, prev_of(trend))
        n = len((node_changes or {}).get(site) or [])
        node_text = None if node_changes is None else (
            f"{n} this month – see the next slide{'s' if n > NODE_ROWS_PER_SLIDE else ''}" if n else "None this month")
        fill_month_exec_slide(prs, slides[ex], cur, prev_of(trend), node_text)
        _set_subtitle(slides[ex], subtitle)
        for s in (slides[dash], slides[ex], *extra[site]):
            fill_month_footer(s, ym)

    cur = total_trend[-1]
    fill_month_dashboard(slides[7], cur, total_trend, prev_of(total_trend))
    node_text = None
    if node_changes is not None:
        node_text = " · ".join(f"{s} {len(node_changes.get(s) or [])}" for s in SITES)
    fill_month_exec_slide(prs, slides[8], cur, prev_of(total_trend), node_text,
                          site_breakdown={s: site_trends[s][-1] for s in SITES})
    _set_subtitle(slides[8], subtitle)
    for s in (slides[7], slides[8]):
        fill_month_footer(s, ym)

    ordered = [slides[0], slides[7], slides[8]]
    for dash, ex, site in positions:
        ordered += [slides[dash], slides[ex]] + extra[site]
    wk.order_slides(prs, ordered)
    prs.save(out_path)
