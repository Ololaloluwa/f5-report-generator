#!/usr/bin/env python3
"""
F5 MONTHLY ACTIVITY REPORT - compiled from the weekly activity reports
=======================================================================

Builds "F5 Monthly Activity Report - MTN Nigeria" (Excel) from the month's
weekly activity workbooks (F5_Weekly_Activity_Report_<dates>.xlsx). Rules
agreed with the user 03-Oct-2026:

  Which weekly files: every file whose Monday-Sunday week touches the month.
    A file's week is read from its Overview title ("(3rd- 9th August)"),
    else from its file name; the year comes from the month being reported.
  Health checks (DNS + LTM): only the month's days are copied, exactly as
    they are in the weekly files (whole day blocks, styles included).
    Days are grouped into "Week N" sheets by calendar week; a leftover of
    1-3 days at the start/end of the month joins the neighbouring full
    week, 4+ days is a week of its own (Aug 2026: 1-9, 10-16, 17-23,
    24-31; Sep 2026: 1-6, 7-13, 14-20, 21-30).
  Cases (Overview table): every case number that appears in any of the
    month's weekly files, once, with its LATEST row (status + response).
    Case numbers are compared without the ** and leading zeros.
  Case screenshot sheets: matched by their "CASE: ..." title (they don't
    show the case number); the latest week's sheet wins.
  Summary: peaks are worked out from the month's daily health-check rows,
    averages too (rounded up to the next 5%: "<60%"); remarks are a
    generated first draft for the engineer to edit.
  Cover page: the cover slide as a picture with the month written on it
    (assets/activity_cover.png + assets/cover_font.ttf). The weekly
    files' embedded PowerPoint cover can't be kept by openpyxl.

Everything that looks wrong but doesn't stop the report is returned as a
warning ([WARN] in the run log), e.g. a day with no weekly file.
"""

import copy
import io
import math
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, PatternFill
from openpyxl.styles.cell_style import StyleArray
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

import generate_weekly_report as wk

ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
COVER_IMAGE = os.path.join(ASSETS, "activity_cover.png")
COVER_FONT = os.path.join(ASSETS, "cover_font.ttf")

DAYS = ("MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY", "SATURDAY", "SUNDAY")
HEALTH_SHEETS = (("DNS", "DNS Health Checks"), ("LTM", "LTM Health Checks"))
MERGE_LEFTOVER_BELOW = 4        # leftover days (start/end of month) fewer than this join the next/previous week
FLAG_PEAK_AT = 85               # % CPU/memory peak that the generated summary line calls out

_MONTHS = ("january", "february", "march", "april", "may", "june", "july", "august",
           "september", "october", "november", "december")


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------

def _ord(d):
    return "th" if 11 <= d % 100 <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(d % 10, "th")


def _short(d):
    return f"{d.day} {d:%b}"


def _span(days):
    a, b = days[0], days[-1]
    if a == b:
        return _short(a)
    return f"{a.day}-{b.day} {b:%b}" if a.month == b.month else f"{_short(a)}-{_short(b)}"


def _month_days(ym):
    d = date(ym[0], ym[1], 1)
    out = []
    while d.month == ym[1]:
        out.append(d)
        d += timedelta(days=1)
    return out


def _month_num(word):
    """'Aug' / 'august' / 'Sept' / 'Septemb' -> 8 / 8 / 9 / 9; None for any other word."""
    w = (word or "").lower().rstrip(".")
    if len(w) < 3:
        return None
    for i, m in enumerate(_MONTHS, start=1):
        if m.startswith(w) or (w.startswith(m[:3]) and len(w) <= len(m) and m.startswith(w[:3])
                               and w in ("sept",)):
            return i
    return None


def pct(v):
    """'< 75%' / '<75%' / 0.75 / 75 -> 75.0 (None if not a percentage)."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) * 100 if v <= 1 else float(v)
    m = re.search(r"(\d+(?:\.\d+)?)", str(v))
    return float(m.group(1)) if m else None


def gbps(v):
    """'12Gbps' / '725Mbps' / '12.5 Gbps' -> Gbps as a float (None if unreadable)."""
    m = re.search(r"(\d+(?:\.\d+)?)\s*([GMK])", str(v or ""), re.I)
    if not m:
        return None
    x = float(m.group(1))
    return {"G": x, "M": x / 1000, "K": x / 1e6}[m.group(2).upper()]


def node_label(name):
    """'NG-ABJ-DNS-WAP-F5BIGIP-01.mtn.com.ng' -> 'ABJ DNS'; 'NG-APP-AIRLB-...' -> 'APP AIRLB'."""
    base = str(name).split(".")[0]
    parts = [p for p in re.split(r"[-_\s]+", base) if p]
    parts = [p for p in parts if p.upper() not in ("NG", "F5", "F5BIGIP") and not p.isdigit()]
    return " ".join(p.upper() for p in parts[:2]) or str(name)


def case_key(text):
    """'**01215073**' / '1209536' -> '1209536' (None if it isn't a case number)."""
    digits = re.sub(r"\D", "", str(text or ""))
    return digits.lstrip("0") or None


def case_sheet_key(title):
    """'CASE: DNS CONFIGURATION ' -> 'DNS CONFIGURATION' (for matching the same case across weeks)."""
    t = re.sub(r"^\s*CASE\s*:?", "", str(title or ""), flags=re.I)
    return re.sub(r"[^A-Z0-9]+", " ", t.upper()).strip() or None


# --------------------------------------------------------------------------
# Which week is each weekly file?
# --------------------------------------------------------------------------

_DATE_TOKEN = re.compile(r"(?<![A-Za-z0-9])(\d{1,2})(?:st|nd|rd|th)?\b\.?\s*([A-Za-z]{3,9})?", re.I)


def parse_week_text(text, ym):
    """
    '(3rd- 9th August)' / '(27th July- 2nd August)' / '...21st_September_-_27th_September'
    -> (monday, sunday) of that week, with the year taken from the report month
    `ym`. None if no two day numbers can be found.
    """
    text = str(text or "").replace("_", " ")
    toks = []
    for m in _DATE_TOKEN.finditer(text):
        day = int(m.group(1))
        mon = _month_num(m.group(2)) if m.group(2) else None
        if 1 <= day <= 31:
            toks.append((day, mon))
    toks = [t for t in toks if t[1] is not None or len(toks) >= 2]
    if len(toks) < 2:
        return None
    (d1, m1), (d2, m2) = toks[0], toks[1]
    if m2 is None:
        m2 = m1
    if m1 is None:
        m1 = m2 if d1 <= d2 else (m2 - 2) % 12 + 1
    if m1 is None or m2 is None:
        return None
    y = ym[0]
    best = None
    for yy in (y - 1, y, y + 1):                 # the year that puts the week nearest the report month
        try:
            end = date(yy, m2, d2)
            start = date(yy - (1 if m1 > m2 else 0), m1, d1)
        except ValueError:
            continue
        dist = abs((end - date(ym[0], ym[1], 15)).days)
        if best is None or dist < best[0]:
            best = (dist, start, end)
    if best is None:
        return None
    monday = wk._week_for_range(best[1], best[2])
    return monday, monday + timedelta(days=6)


@dataclass
class WeeklyFile:
    key: tuple                  # (monday, sunday)
    path: str
    name: str                   # file name as uploaded (for messages)
    wb: object = None


def load_weekly_files(paths, ym, warn, names=None):
    """Opens every weekly activity file that touches month `ym` -> {week key: WeeklyFile}."""
    first, last = _month_days(ym)[0], _month_days(ym)[-1]
    found = {}
    for i, path in enumerate(paths):
        name = (names[i] if names else None) or os.path.basename(path)
        # a cheap first look (title only) so files for other weeks are never fully opened
        try:
            peek = openpyxl.load_workbook(path, read_only=True)
            title = None
            if "Overview" in peek.sheetnames:
                title = next(peek["Overview"].iter_rows(min_row=1, max_row=1, max_col=1, values_only=True))[0]
            peek.close()
        except Exception as e:                   # noqa: BLE001 - one bad file mustn't stop the rest
            warn(f"Activity report '{name}' couldn't be opened ({e}) - skipped")
            continue
        key = parse_week_text(title, ym) or parse_week_text(name, ym)
        if key is None:
            warn(f"Activity report '{name}': couldn't tell which week it is from its title "
                 f"('{str(title or '').strip()[:60]}') or its file name - skipped")
            continue
        if key[1] < first or key[0] > last:
            continue                             # a week of another month - not needed
        if key in found:
            warn(f"Two activity reports for the week {_short(key[0])}-{_short(key[1])}: "
                 f"'{found[key].name}' and '{name}' - used '{name}'")
        found[key] = WeeklyFile(key=key, path=path, name=name)
    for f in found.values():
        f.wb = openpyxl.load_workbook(f.path)
    return found


# --------------------------------------------------------------------------
# Copying cells / sheets between workbooks
# --------------------------------------------------------------------------

def _copy_style(src, dst):
    if src.has_style:
        dst.font = copy.copy(src.font)
        dst.fill = copy.copy(src.fill)
        dst.border = copy.copy(src.border)
        dst.alignment = copy.copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy.copy(src.protection)


def copy_rows(src, dst, r0, r1, dst_r0, with_columns=False, with_images=True):
    """
    Copies rows r0..r1 of worksheet `src` to `dst` starting at row dst_r0:
    values, styles, merged cells, row heights, and pictures anchored in
    those rows. Only cells that exist are visited (some sheets claim to be
    16,384 columns wide).
    """
    off = dst_r0 - r0
    for rng in src.merged_cells.ranges:
        if rng.min_row >= r0 and rng.max_row <= r1:
            dst.merge_cells(start_row=rng.min_row + off, start_column=rng.min_col,
                            end_row=rng.max_row + off, end_column=rng.max_col)
    for (r, c), cell in list(src._cells.items()):
        if r < r0 or r > r1:
            continue
        out = dst.cell(row=r + off, column=c)
        if not isinstance(cell, MergedCell) and not isinstance(out, MergedCell):
            out.value = cell.value
        _copy_style(cell, out)
    for r in range(r0, r1 + 1):
        dim = src.row_dimensions.get(r)
        if dim is not None and dim.height:
            dst.row_dimensions[r + off].height = dim.height
    if with_columns:
        for key, dim in src.column_dimensions.items():
            nd = dst.column_dimensions[key]
            nd.width, nd.hidden = dim.width, dim.hidden
            nd.min, nd.max = dim.min, dim.max
        dst.sheet_view.showGridLines = src.sheet_view.showGridLines
        if src.sheet_view.zoomScale:
            dst.sheet_view.zoomScale = src.sheet_view.zoomScale
    if with_images:
        for img in getattr(src, "_images", []):
            row = img.anchor._from.row + 1
            if r0 <= row <= r1:
                new = XLImage(io.BytesIO(img._data()))
                new.width, new.height = img.width, img.height
                anchor = copy.deepcopy(img.anchor)
                anchor._from.row += off
                if getattr(anchor, "to", None) is not None:
                    anchor.to.row += off
                new.anchor = anchor
                dst.add_image(new)


def _last_row(ws):
    return max((r for (r, _c), cell in ws._cells.items() if cell.value is not None), default=1)


# --------------------------------------------------------------------------
# Health checks
# --------------------------------------------------------------------------

def day_blocks(ws):
    """{day index 0-6: (first row, last row)} for a weekly health-check sheet, plus the header's last row."""
    heads = []
    for (r, c), cell in ws._cells.items():
        if c == 1 and isinstance(cell.value, str) and cell.value.strip().upper() in DAYS:
            heads.append((r, DAYS.index(cell.value.strip().upper())))
    heads.sort()
    blocks = {}
    last = _last_row(ws)
    for i, (r, d) in enumerate(heads):
        end = heads[i + 1][0] - 1 if i + 1 < len(heads) else last
        blocks.setdefault(d, (r, end))
    header_end = heads[0][0] - 1 if heads else 0
    return blocks, header_end


def _cell_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 40000 < v < 60000:          # an Excel date serial shown as a number
        return date(1899, 12, 30) + timedelta(days=int(v))
    m = re.match(r"^\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*$", str(v or ""))
    if m:
        a, b, y = (int(x) for x in m.groups())
        for mo, d in ((a, b), (b, a)):         # m/d/yyyy is what the files use; d/m as a fallback
            try:
                return date(y, mo, d)
            except ValueError:
                pass
    return None


def group_weeks(ym):
    """The month's days as [[dates of Week 1], [Week 2], ...] (leftover rule in the module docstring)."""
    chunks = []
    for d in _month_days(ym):
        if chunks and wk.monday_of(chunks[-1][0]) == wk.monday_of(d):
            chunks[-1].append(d)
        else:
            chunks.append([d])
    if len(chunks) > 1 and len(chunks[0]) < MERGE_LEFTOVER_BELOW:
        chunks[1] = chunks[0] + chunks[1]
        chunks.pop(0)
    if len(chunks) > 1 and len(chunks[-1]) < MERGE_LEFTOVER_BELOW:
        chunks[-2] = chunks[-2] + chunks[-1]
        chunks.pop()
    return chunks


@dataclass
class HealthStats:
    nodes: set = field(default_factory=set)
    cpu: list = field(default_factory=list)
    mem: list = field(default_factory=list)
    peak: dict = field(default_factory=dict)   # {"cpu"/"mem"/"in"/"out": (value, raw cell value, node, date)}


def _scan_rows(ws, r0, r1, day, stats):
    cols = {}
    for r in range(r0, r1 + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.strip().upper() == "S/N":
            for c in range(1, 16):
                h = str(ws.cell(r, c).value or "").lower()
                if "node name" in h:
                    cols["node"] = c
                elif h.startswith("cpu"):
                    cols["cpu"] = c
                elif h.startswith("memory"):
                    cols["mem"] = c
                elif "throughput in" in h:
                    cols["in"] = c
                elif "throughput out" in h:
                    cols["out"] = c
            continue
        if "node" not in cols:
            continue
        node = ws.cell(r, cols["node"]).value
        if not isinstance(node, str) or not node.strip():
            continue
        node = node.strip()
        stats.nodes.add(node)
        for key, conv in (("cpu", pct), ("mem", pct), ("in", gbps), ("out", gbps)):
            if key not in cols:
                continue
            raw = ws.cell(r, cols[key]).value
            val = conv(raw)
            if val is None:
                continue
            if key in ("cpu", "mem"):
                getattr(stats, key).append(val)
            if key not in stats.peak or val > stats.peak[key][0]:
                stats.peak[key] = (val, raw, node, day)


def build_health_sheets(out_wb, files, ym, warn):
    """Adds the DNS / LTM 'Week N' sheets; returns {"DNS": HealthStats, "LTM": HealthStats}."""
    stats = {kind: HealthStats() for kind, _ in HEALTH_SHEETS}
    missing_days = []
    groups = group_weeks(ym)
    for kind, sheet_name in HEALTH_SHEETS:
        for n, days in enumerate(groups, start=1):
            ws_out = None
            row = 1
            for d in days:
                key = (wk.monday_of(d), wk.monday_of(d) + timedelta(days=6))
                f = files.get(key)
                if f is None:
                    if kind == "DNS":
                        missing_days.append(d)
                    continue
                src = next((f.wb[s] for s in f.wb.sheetnames if s.strip().lower() == sheet_name.lower()), None)
                if src is None:
                    warn(f"'{f.name}' has no '{sheet_name}' sheet - {_short(d)} left out of the {kind} checks")
                    continue
                blocks, header_end = day_blocks(src)
                block = blocks.get(d.weekday())
                if block is None:
                    warn(f"'{f.name}' {sheet_name}: no {DAYS[d.weekday()].title()} section - "
                         f"{_short(d)} left out")
                    continue
                if ws_out is None:
                    ws_out = out_wb.create_sheet(f"{sheet_name} Week {n}")
                    if header_end:
                        copy_rows(src, ws_out, 1, header_end, 1, with_columns=True)
                    else:
                        copy_rows(src, ws_out, 1, 0, 1, with_columns=True)
                    row = header_end + 1
                r0, r1 = block
                copy_rows(src, ws_out, r0, r1, row)
                for r in range(row, row + r1 - r0 + 1):        # date serials typed into a General cell -> show as dates
                    c = ws_out.cell(r, 2)
                    if isinstance(c.value, (int, float)) and 40000 < c.value < 60000 and \
                            c.number_format == "General":
                        c.number_format = "m/d/yyyy"
                _scan_rows(src, r0, r1, d, stats[kind])
                wrong = sorted({cd for r in range(r0, r1 + 1)
                                if (cd := _cell_date(src.cell(r, 2).value)) and cd != d})
                if wrong:
                    warn(f"{kind} health checks, {DAYS[d.weekday()].title()} {_short(d)} (from '{f.name}'): "
                         f"the Date column says {', '.join(_short(x) for x in wrong)} - copied as it is, check it")
                row += r1 - r0 + 1
    if missing_days:
        days = sorted(set(missing_days))
        weeks = sorted({wk.monday_of(d) for d in days})
        need = ", ".join(f"{_short(m)}-{_short(m + timedelta(days=6))}" for m in weeks)
        warn(f"No weekly activity report for {need} - {len(days)} day(s) of {days[0]:%B} "
             f"({', '.join(_span([d for d in days if wk.monday_of(d) == m]) for m in weeks)}) "
             f"are missing from the health checks")
    return stats, groups


# --------------------------------------------------------------------------
# Cases
# --------------------------------------------------------------------------

@dataclass
class CaseRow:
    key: str
    values: dict                 # {column number: value} for A, B, C, E, G, I
    height: float = None
    week: tuple = None


def read_case_rows(ws):
    """The CASE NO table on a weekly Overview sheet -> [CaseRow] (in order)."""
    header = None
    for (r, c), cell in ws._cells.items():
        if c == 1 and isinstance(cell.value, str) and cell.value.strip().upper() == "CASE NO":
            header = r if header is None else min(header, r)
    if header is None:
        return [], None
    rows = []
    r = header + 1
    while r <= _last_row(ws):
        key = case_key(ws.cell(r, 1).value)
        if key is None:
            if any(ws.cell(r, c).value for c in (2, 3, 5)):
                r += 1
                continue
            break
        vals = {c: ws.cell(r, c).value for c in (1, 2, 3, 5, 7, 9)}
        dim = ws.row_dimensions.get(r)
        rows.append(CaseRow(key=key, values=vals, height=dim.height if dim else None))
        r += 1
    return rows, header


def merge_cases(files):
    """Every case in the month's weekly files, once, latest week's row winning; first-seen order."""
    order, latest = [], {}
    for key in sorted(files):
        ov = files[key].wb["Overview"] if "Overview" in files[key].wb.sheetnames else None
        if ov is None:
            continue
        rows, _ = read_case_rows(ov)
        for row in rows:
            row.week = key
            if row.key not in latest:
                order.append(row.key)
            latest[row.key] = row
    return [latest[k] for k in order]


def merge_case_sheets(files):
    """{normalized case title: (file, sheet)} latest week winning, in first-seen order."""
    order, latest = [], {}
    for key in sorted(files):
        f = files[key]
        for name in f.wb.sheetnames:
            if "case" not in name.lower():
                continue
            ws = f.wb[name]
            k = case_sheet_key(ws["H1"].value) or case_sheet_key(name)
            if k not in latest:
                order.append(k)
            latest[k] = (f, ws)
    return [latest[k] for k in order]


def _row_height(texts_widths):
    lines = 1
    for text, width in texts_widths:
        if not text:
            continue
        chars = max(1, int(width * 1.0))
        n = sum(max(1, math.ceil(len(part) / chars)) for part in str(text).split("\n"))
        lines = max(lines, n)
    return max(15.0, lines * 14.5 + 4)


def write_overview(ws, ym, cases, stats):
    days = _month_days(ym)
    a, b = days[0], days[-1]
    ws["A1"].value = (f"F5 Monthly Activity Report - MTN Nigeria\n"
                      f"({a.day}{_ord(a.day)} {a:%B} - {b.day}{_ord(b.day)} {b:%B} {b.year})")
    for (r, c), cell in list(ws._cells.items()):
        if isinstance(cell, MergedCell) or not isinstance(cell.value, str):
            continue
        v = cell.value
        if r <= 6 and re.match(r"^\s*DNS\s*\(\d+\)\s*$", v) and ws.cell(r, 1).value and \
                "health" in str(ws.cell(r, 1).value).lower():
            cell.value = f"  DNS ({len(stats['DNS'].nodes)})"
        elif r <= 6 and re.match(r"^\s*LTM\s*\(\d+\)\s*$", v) and ws.cell(r, 1).value and \
                "health" in str(ws.cell(r, 1).value).lower():
            cell.value = f"LTM ({len(stats['LTM'].nodes)})"
        elif r <= 6 and "(Monday - Sunday)" in v:
            cell.value = f" {len(days)}   (Monday - Sunday)"

    _, header = read_case_rows(ws)
    if header is None:
        return
    first = header + 1
    width = {c: (ws.column_dimensions[get_column_letter(c)].width or 9) for c in range(1, 11)}
    styles = {c: copy.copy(ws.cell(first, c)._style) for c in range(1, 11)}
    old_last = max(_last_row(ws), first)
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= first:
            ws.unmerge_cells(str(rng))
    for r in range(first, old_last + 1):
        for c in range(1, 11):
            ws.cell(r, c).value = None
            ws.cell(r, c)._style = copy.copy(styles[c])
        ws.row_dimensions[r].height = None
    for i, case in enumerate(cases):
        r = first + i
        for c in range(1, 11):
            ws.cell(r, c)._style = copy.copy(styles[c])
        for c, v in case.values.items():
            ws.cell(r, c).value = v
        for c0, c1 in ((3, 4), (5, 6), (7, 8), (9, 10)):
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
        need = _row_height([(case.values.get(2), width[2]),
                            (case.values.get(7), width[7] + width[8]),
                            (case.values.get(9), width[9] + width[10])])
        ws.row_dimensions[r].height = max(need, case.height or 0)
    for r in range(first + len(cases), old_last + 1):       # rows the old table used but this one doesn't
        for c in range(1, 11):
            ws.cell(r, c)._style = StyleArray()


# --------------------------------------------------------------------------
# Summary
# --------------------------------------------------------------------------

def _avg_text(values):
    if not values:
        return "n/a"
    avg = sum(values) / len(values)
    return f"<{min(100, (math.floor(avg / 5) + 1) * 5)}%"          # 55.9 -> <60%, 60.0 -> <65%


RED, AMBER, GREEN = "FFFF0000", "FFFFFF00", "FF00B050"


def _level_fill(value):
    """Red at FLAG_PEAK_AT% and over, yellow from 75%, green below - the colours the weekly summaries use."""
    colour = RED if value >= FLAG_PEAK_AT else AMBER if value >= 75 else GREEN
    return PatternFill(fill_type="solid", start_color=colour, end_color=colour)


def _and(items):
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


def write_summary(ws, stats, ym):
    """Fills the Monthly DNS / LTM highlight rows and the summary line, keeping the sheet's layout."""
    section = None
    flagged = []
    labels = {"cpu": "CPU", "mem": "memory"}
    for kind in ("DNS", "LTM"):
        for key in ("cpu", "mem"):
            p = stats[kind].peak.get(key)
            if p and p[0] >= FLAG_PEAK_AT:
                flagged.append(f"{node_label(p[2])} ({labels[key]} peaked at {p[0]:.0f}%)")
    for row in ws.iter_rows():
        cell_b = row[1] if len(row) > 1 else None
        if cell_b is None or isinstance(cell_b, MergedCell) or not isinstance(cell_b.value, str):
            continue
        text = cell_b.value.strip()
        low = text.lower()
        r = cell_b.row
        val = ws.cell(r, 3)
        if "health check highlights" in low:
            section = "DNS" if "dns" in low else "LTM"
            cell_b.value = f"Monthly {section} Health Check Highlights"
            continue
        if low == "summary":
            nxt = ws.cell(r + 1, 2)
            if not isinstance(nxt, MergedCell):
                if flagged:
                    nxt.value = (" All DNS and LTM devices operated within safe thresholds, except "
                                 + _and(flagged) + ".")
                else:
                    nxt.value = " All DNS and LTM devices operated within safe thresholds."
            section = None
            continue
        if section is None or isinstance(val, MergedCell):
            continue
        st = stats[section]
        if low.startswith("total devices"):
            val.value = len(st.nodes)
        elif low.startswith("average cpu"):
            val.value = _avg_text(st.cpu)
        elif low.startswith("average memory"):
            val.value = _avg_text(st.mem)
        elif low.startswith(("peak cpu", "peak memory", "peak throughput in", "peak throughput out")):
            key = ("cpu" if "cpu" in low else "mem" if "memory" in low
                   else "in" if "throughput in" in low else "out")
            p = st.peak.get(key)
            base = re.sub(r"\s*\(.*$", "", text)
            if p is None:
                cell_b.value, val.value = base, "n/a"
                continue
            cell_b.value = f"{base} ({node_label(p[2])})"
            if key in ("cpu", "mem"):
                val.value = p[0] / 100
                val.number_format = "0%"
                val.fill = _level_fill(p[0])
            else:
                val.value = str(p[1]).strip()
        elif low.startswith("remarks"):
            parts = []
            c, m = st.peak.get("cpu"), st.peak.get("mem")
            i, o = st.peak.get("in"), st.peak.get("out")
            if c:
                parts.append(f"Highest CPU: {node_label(c[2])} at {c[0]:.0f}% ({c[3]:%a} {_short(c[3])}).")
            if m:
                parts.append(f"Highest memory: {node_label(m[2])} at {m[0]:.0f}% ({m[3]:%a} {_short(m[3])}).")
            if i and o:
                parts.append(f"Peak throughput: {str(i[1]).strip()} in ({node_label(i[2])}), "
                             f"{str(o[1]).strip()} out ({node_label(o[2])}).")
            val.value = " ".join(parts) or "No readings found for this month."
            worst = max((st.peak[k][0] for k in ("cpu", "mem") if k in st.peak), default=0)
            val.fill = _level_fill(worst)
            val.alignment = Alignment(horizontal=val.alignment.horizontal,
                                      vertical=val.alignment.vertical or "top", wrap_text=True)


# --------------------------------------------------------------------------
# Cover page
# --------------------------------------------------------------------------

def add_cover(out_wb, ym, warn):
    ws = out_wb.create_sheet("Cover page", 0)
    ws.sheet_view.showGridLines = False
    if not os.path.exists(COVER_IMAGE):
        ws["B2"].value = f"F5 MONTHLY REPORT - {date(ym[0], ym[1], 1):%B %Y}".upper()
        warn("Cover picture (assets/activity_cover.png) not found - a text cover was used")
        return
    from PIL import Image, ImageDraw, ImageFont
    im = Image.open(COVER_IMAGE).convert("RGB")
    w, h = im.size
    try:
        font = ImageFont.truetype(COVER_FONT, int(h * 0.0345))
    except OSError:
        font = ImageFont.load_default(size=int(h * 0.0345))
    ImageDraw.Draw(im).text((w / 2, h * 0.704), f"{date(ym[0], ym[1], 1):%B %Y}".upper(),
                            font=font, fill=(0, 0, 0), anchor="mm")
    buf = io.BytesIO()
    im.save(buf, format="PNG", optimize=True)
    buf.seek(0)
    img = XLImage(buf)
    img.width, img.height = 960, int(960 * h / w)
    ws.add_image(img, "B2")


# --------------------------------------------------------------------------
# The whole workbook
# --------------------------------------------------------------------------

@dataclass
class ActivityResult:
    path: str
    weeks_used: list                 # [(week key, file name)]
    cases: int
    case_sheets: int
    health_weeks: list               # [[dates], ...]
    stats: dict


def build_activity_report(paths, ym, out_path, warn=print, names=None):
    """
    paths: the weekly activity workbooks (any order; extra weeks are ignored).
    Returns ActivityResult, or None if none of the files belongs to the month.
    """
    files = load_weekly_files(paths, ym, warn, names=names)
    if not files:
        warn(f"None of the {len(paths)} activity report(s) is for {date(ym[0], ym[1], 1):%B %Y} - "
             f"no activity report built")
        return None
    latest = files[max(files)]
    # Start from the latest weekly file so the workbook keeps its theme and the
    # Overview / Summary layout; everything else is rebuilt.
    out = openpyxl.load_workbook(latest.path)
    for name in list(out.sheetnames):
        if name not in ("Overview", "Summary"):
            del out[name]

    stats, groups = build_health_sheets(out, files, ym, warn)
    cases = merge_cases(files)
    if "Overview" in out.sheetnames:
        write_overview(out["Overview"], ym, cases, stats)
    sheets = merge_case_sheets(files)
    for i, (f, src) in enumerate(sheets, start=1):
        dst = out.create_sheet(f"Cases Logged {i}")
        copy_rows(src, dst, 1, max(_last_row(src), max((im.anchor._from.row + 1 for im in src._images), default=1)),
                  1, with_columns=True)
    if "Summary" in out.sheetnames:
        write_summary(out["Summary"], stats, ym)
        out.move_sheet("Summary", offset=len(out.sheetnames) - 1 - out.sheetnames.index("Summary"))
    add_cover(out, ym, warn)
    if "Overview" in out.sheetnames:
        out.move_sheet("Overview", offset=1 - out.sheetnames.index("Overview"))
    out.active = 0
    for ws in out.worksheets:
        ws.sheet_view.tabSelected = ws is out.worksheets[0]
    out.save(out_path)
    for f in files.values():
        f.wb.close()
    return ActivityResult(
        path=out_path, weeks_used=[(k, files[k].name) for k in sorted(files)], cases=len(cases),
        case_sheets=len(sheets), health_weeks=groups, stats=stats,
    )


def activity_file_name(ym):
    return f"F5_Monthly_Activity_Report_MTN_Nigeria_{date(ym[0], ym[1], 1):%B}{ym[0]}.xlsx"
