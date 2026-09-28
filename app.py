"""
MTN F5 Report Generator - web app
===================================

A small Flask site in front of the SAME code the .bat file runs
(run_weekly_reports.generate_reports / validate_outputs.run_checks) - it
never re-implements report logic, so the website and the .bat file can't
drift apart.

Flow:
  1. /login        shared team password (APP_PASSWORD environment variable)
  2. /             choose report type (Weekly; Monthly "coming soon") and
                   upload MTN_F5_ATTACKS_SOURCE.xlsx. The saved default
                   Certificate.xlsx is used unless they upload a different one.
  3. /job/<id>     sheet check + week dropdown + next week's sheet names
  4. generate      builds the deck + Excel, runs validation, shows results
                   (a [FAIL] shows a red warning banner but downloads stay
                   available - confirmed with the user 28-Sep-2026)
  5. download      the two files

Nothing is kept: every upload lives in its own temporary folder that is
deleted when the user clicks "Done", or automatically JOB_TTL_MINUTES
after the upload - whichever comes first.

This f5-report-generator/ folder is self-contained - it is the git repository that gets
deployed: the report code (generate_weekly_report.py etc.), both blank
templates, requirements.txt and render.yaml all live here. The .bat
version in the folder above runs this same code (it calls
f5-report-generator\run_weekly_reports.py), so there is exactly one copy of everything.

Default certificate list - the first of these that exists:
  1. CERTS_FILE environment variable (a path)
  2. /etc/secrets/Certificate.xlsx.b64 - a Render "Secret File" holding the
                                       workbook as base64 text (Render's secret
                                       files are pasted in as text - see
                                       README.md), so the MTN certificate list
                                       never goes into GitHub
     (/etc/secrets/Certificate.xlsx is also accepted, if uploaded as-is)
  3. ./Certificate.xlsx or ../Certificate.xlsx - local runs only (the one
                                       the .bat file uses; never committed)

Environment variables:
  APP_PASSWORD  (required) the shared team password
  SECRET_KEY    (recommended) signs the login cookie; a random one is used
                if missing, which just means everyone is logged out on restart
  JOBS_DIR      (optional) where temporary job folders go; default: system temp
"""

import base64
import hmac
import os
import secrets
import shutil
import sys
import tempfile
import threading
import time
from datetime import date, timedelta
from functools import wraps

from flask import (
    Flask, abort, flash, redirect, render_template, request, send_file,
    session, url_for,
)

# The report code sits right next to this file.
APP_DIR = os.path.dirname(os.path.abspath(__file__))
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)

import openpyxl                               # noqa: E402
import generate_weekly_report as pptx_gen     # noqa: E402
import run_weekly_reports                     # noqa: E402
import validate_outputs                       # noqa: E402

PPTX_TEMPLATE = os.path.join(APP_DIR, "MTN_Security_Metrics_Report_SOURCE.pptx")
XLSX_TEMPLATE = os.path.join(APP_DIR, "MTN_F5_WAF_Weekly_Report_Temp.xlsx")

JOB_TTL_MINUTES = 15
MAX_UPLOAD_MB = 250          # the source workbook grows ~2-3 MB a week
PROPRIETARY_NOTICE = (
    "This tool is the proprietary property of Olola. It may not be used without "
    "payment/authorization from Olola. This tool is provided as-is. Olola accepts no "
    "liability for any errors, data loss, or other issues arising from its use."
)

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ.get("SECRET_KEY") or secrets.token_hex(32),
    MAX_CONTENT_LENGTH=MAX_UPLOAD_MB * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("RENDER") is not None,   # HTTPS on Render
    PERMANENT_SESSION_LIFETIME=timedelta(hours=8),
)

JOBS_DIR = os.environ.get("JOBS_DIR") or os.path.join(tempfile.gettempdir(), "f5_report_jobs")
os.makedirs(JOBS_DIR, exist_ok=True)

# The free Render plan has 512 MB of memory and one report run peaks around
# 190 MB, so only one report is generated at a time; anyone else simply
# waits a few seconds for theirs.
_generate_lock = threading.Lock()


def _decoded_secret_certs():
    """Decode the base64 Render Secret File to a real .xlsx (re-done if the secret changes)."""
    src = "/etc/secrets/Certificate.xlsx.b64"
    if not os.path.isfile(src):
        return None
    out = os.path.join(JOBS_DIR, "_default_Certificate.xlsx")
    if not os.path.exists(out) or os.path.getmtime(out) < os.path.getmtime(src):
        try:
            with open(src, "rb") as f:
                data = base64.b64decode(b"".join(f.read().split()))
            with open(out, "wb") as f:
                f.write(data)
        except (ValueError, OSError):
            app.logger.exception("couldn't decode %s", src)
            return None
    return out


def default_certs_path():
    """The saved default Certificate.xlsx, or None if there isn't one."""
    for path in (os.environ.get("CERTS_FILE"),
                 _decoded_secret_certs(),
                 "/etc/secrets/Certificate.xlsx",
                 os.path.join(APP_DIR, "Certificate.xlsx"),
                 os.path.join(os.path.dirname(APP_DIR), "Certificate.xlsx")):
        if path and os.path.isfile(path):
            return path
    return None


_certs_info_cache = {}


def certs_info(path):
    """
    'Checked 09 Sep 2026 by Aniebiet Joseph' for a certificate workbook -
    shown on the page so people can tell when the saved default is getting
    old. Cached per file version; None if it can't be read.
    """
    if not path:
        return None
    key = (path, os.path.getmtime(path))
    if key not in _certs_info_cache:
        try:
            data = pptx_gen.read_certificates(path, date.today().year)
            first = next(iter(data.values()), {})
            checked = first.get("checked")
            text = f"checked {checked:%d %b %Y}" if isinstance(checked, date) else "check date not filled in"
            if first.get("engineer"):
                text += f" by {first['engineer']}"
            _certs_info_cache[key] = text
        except Exception:
            _certs_info_cache[key] = None
    return _certs_info_cache[key]


@app.context_processor
def _globals():
    return {"notice": PROPRIETARY_NOTICE, "ttl": JOB_TTL_MINUTES}


# --------------------------------------------------------------------------
# Login (one shared team password)
# --------------------------------------------------------------------------

def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("auth"):
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)
    return wrapped


@app.route("/login", methods=["GET", "POST"])
def login():
    password = os.environ.get("APP_PASSWORD")
    if not password:
        return render_template("message.html", title="Not set up yet",
                               message="The APP_PASSWORD setting hasn't been set on the server, "
                                       "so nobody can log in yet."), 503
    if request.method == "POST":
        if hmac.compare_digest(request.form.get("password", "").encode(), password.encode()):
            session.clear()
            session.permanent = True
            session["auth"] = True
            nxt = request.args.get("next", "")
            return redirect(nxt if nxt.startswith("/") and not nxt.startswith("//") else url_for("index"))
        time.sleep(1)   # slow down password guessing
        flash("Wrong password.", "error")
    return render_template("login.html")


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login"))


# --------------------------------------------------------------------------
# Temporary job folders
# --------------------------------------------------------------------------

def _job_dir(job_id):
    """Folder for one upload. The id is a random hex token; anything else is rejected."""
    if not job_id or len(job_id) != 32 or any(c not in "0123456789abcdef" for c in job_id):
        abort(404)
    path = os.path.join(JOBS_DIR, job_id)
    if not os.path.isdir(path):
        abort(404)
    return path


def _sweep_old_jobs():
    """Delete every job folder older than JOB_TTL_MINUTES."""
    cutoff = time.time() - JOB_TTL_MINUTES * 60
    for name in os.listdir(JOBS_DIR):
        path = os.path.join(JOBS_DIR, name)
        try:
            if os.path.isdir(path) and os.path.getmtime(os.path.join(path, "created")) < cutoff:
                shutil.rmtree(path, ignore_errors=True)
        except OSError:
            shutil.rmtree(path, ignore_errors=True)


def _sweeper():
    while True:
        time.sleep(60)
        _sweep_old_jobs()


threading.Thread(target=_sweeper, daemon=True).start()


def _read_weeks(source_path):
    """
    Sheet NAMES only (about a second, even for the full 29 MB workbook) ->
    the weeks for the dropdown, the default week, the sheet check and
    next week's names.
    """
    wb = openpyxl.load_workbook(source_path, read_only=True)
    try:
        names = list(wb.sheetnames)
    finally:
        wb.close()

    class _Names:
        sheetnames = names
    weeks = pptx_gen.group_sheets_by_week(_Names)
    if not weeks:
        return None
    default = pptx_gen.pick_target_week(weeks)
    latest = pptx_gen.latest_real_week(weeks)
    recent = sorted(weeks, key=lambda k: k[1])
    focus = recent[max(0, recent.index(default) - 1): recent.index(default) + 1]
    options = []
    for key in sorted(weeks, key=lambda k: k[1], reverse=True):
        start, end = key
        missing = [s for s in pptx_gen.SITES if s not in weeks[key]]
        label = f"{start:%d %b} - {end:%d %b %Y}"
        label += f"  (missing {', '.join(missing)})" if missing else "  (IKY, VGC, OJT)"
        if start > date.today():
            label += "  - dated in the future?"
        options.append({"value": end.isoformat(), "label": label, "selected": key == default})
    return {
        "options": options,
        "sheet_warnings": pptx_gen.check_sheet_names(names, focus_weeks=focus),
        "next_names": pptx_gen.suggest_sheet_names(latest[1])[1],
    }


# --------------------------------------------------------------------------
# Pages
# --------------------------------------------------------------------------

@app.route("/")
@login_required
def index():
    path = default_certs_path()
    return render_template("index.html", default_certs=path is not None,
                           default_certs_info=certs_info(path))


@app.route("/upload", methods=["POST"])
@login_required
def upload():
    if request.form.get("report_type", "weekly") != "weekly":
        flash("Only the weekly report is available so far.", "error")
        return redirect(url_for("index"))

    source = request.files.get("source")
    certs = request.files.get("certs")
    if not source or not source.filename:
        flash("Please choose the source spreadsheet (MTN_F5_ATTACKS_SOURCE.xlsx).", "error")
        return redirect(url_for("index"))
    for f in (source, certs):
        if f and f.filename and not f.filename.lower().endswith(".xlsx"):
            flash(f"'{f.filename}' isn't an .xlsx file.", "error")
            return redirect(url_for("index"))

    _sweep_old_jobs()
    job_id = secrets.token_hex(16)
    path = os.path.join(JOBS_DIR, job_id)
    os.makedirs(path)
    open(os.path.join(path, "created"), "w").close()
    source.save(os.path.join(path, "source.xlsx"))
    if certs and certs.filename:
        certs.save(os.path.join(path, "certs.xlsx"))

    try:
        info = _read_weeks(os.path.join(path, "source.xlsx"))
    except Exception:
        shutil.rmtree(path, ignore_errors=True)
        flash("That file couldn't be opened as an Excel workbook.", "error")
        return redirect(url_for("index"))
    if info is None:
        shutil.rmtree(path, ignore_errors=True)
        flash("No weekly sheets were found in that workbook - sheet names should look like "
              "'14th - 20th Sep 2026 for IKY'.", "error")
        return redirect(url_for("index"))
    return redirect(url_for("job", job_id=job_id))


@app.route("/job/<job_id>")
@login_required
def job(job_id):
    path = _job_dir(job_id)
    info = _read_weeks(os.path.join(path, "source.xlsx"))
    return render_template("week.html", job_id=job_id, info=info, step=2,
                           certs=_certs_choice(path))


def _certs_choice(job_path):
    """(path, description) of the certificate file this job will use, or (None, reason)."""
    uploaded = os.path.join(job_path, "certs.xlsx")
    if os.path.exists(uploaded):
        info = certs_info(uploaded)
        return uploaded, "the certificate file you uploaded" + (f" ({info})" if info else "")
    default = default_certs_path()
    if default:
        info = certs_info(default)
        return default, "the saved default certificate list" + (f" ({info})" if info else "")
    return None, "no certificate list - there's no saved default and none was uploaded, " \
                 "so the certificate slides will be left out"


@app.route("/job/<job_id>/generate", methods=["POST"])
@login_required
def generate(job_id):
    path = _job_dir(job_id)
    try:
        week_end = date.fromisoformat(request.form["week_end"])
    except (KeyError, ValueError):
        abort(400)

    out_dir = os.path.join(path, "output")
    shutil.rmtree(out_dir, ignore_errors=True)   # re-generating another week replaces the last one
    with _generate_lock:
        try:
            result = run_weekly_reports.generate_reports(
                os.path.join(path, "source.xlsx"), PPTX_TEMPLATE, XLSX_TEMPLATE, out_dir,
                week_end=week_end, log=lambda _m: None, certs=_certs_choice(path)[0])
            checks = validate_outputs.run_checks(
                os.path.join(path, "source.xlsx"), week_end, log=lambda _m: None)
        except Exception as e:
            app.logger.exception("report generation failed")
            return render_template("message.html", title="Something went wrong",
                                   message=f"The report couldn't be generated: {e}. Nothing was "
                                           f"produced - check the source spreadsheet, or send this "
                                           f"message to Olola.",
                                   back=url_for("job", job_id=job_id)), 500

    lines = [l for l in checks.lines if l.startswith("[")]
    warnings = list(dict.fromkeys(result.warnings + [l for l in lines if l.startswith("[WARN]")]))
    return render_template(
        "result.html", job_id=job_id, result=result, ok=checks.ok,
        fails=[l for l in lines if l.startswith("[FAIL]")], warnings=warnings,
        passes=[l for l in lines if l.startswith("[PASS]")],
        pptx_name=os.path.basename(result.pptx_path), xlsx_name=os.path.basename(result.xlsx_path),
        certs=_certs_choice(path), step=3)


@app.route("/job/<job_id>/download/<kind>")
@login_required
def download(job_id, kind):
    path = _job_dir(job_id)
    out_dir = os.path.join(path, "output")
    wanted = {"pptx": ".pptx", "xlsx": ".xlsx"}.get(kind)
    if wanted is None or not os.path.isdir(out_dir):
        abort(404)
    files = [f for f in os.listdir(out_dir) if f.endswith(wanted)]
    if not files:
        abort(404)
    return send_file(os.path.join(out_dir, files[0]), as_attachment=True, download_name=files[0])


@app.route("/job/<job_id>/done", methods=["POST"])
@login_required
def done(job_id):
    shutil.rmtree(_job_dir(job_id), ignore_errors=True)
    flash("Your files have been deleted from the server.", "ok")
    return redirect(url_for("index"))


@app.errorhandler(413)
def too_big(_e):
    flash(f"That file is bigger than {MAX_UPLOAD_MB} MB.", "error")
    return redirect(url_for("index"))


@app.errorhandler(404)
def not_found(_e):
    return render_template("message.html", title="Not found",
                           message=f"That upload doesn't exist any more - files are deleted "
                                   f"{JOB_TTL_MINUTES} minutes after uploading. Please upload again.",
                           back=url_for("index")), 404


if __name__ == "__main__":
    # Local testing only (see README.md). On Render, gunicorn runs the app.
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", 5000)), debug=False)
