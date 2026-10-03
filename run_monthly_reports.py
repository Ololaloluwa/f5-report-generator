#!/usr/bin/env python3
"""
Monthly MTN F5 Reports - run everything
==========================================

One command that builds the monthly outputs from MTN_F5_ATTACKS_SOURCE.xlsx
and then checks them:
  1. PowerPoint deck      (generate_monthly_report.py, weekly template)
  2. Detailed WAF Excel   (generate_monthly_excel.py, weekly template)
  2b. Monthly activity report (generate_activity_report.py) - only when the
      month's weekly activity reports are given (--activity-dir)
  3. Checks               ([PASS] / [WARN] / [FAIL] lines - a [FAIL] means
                           don't send it without looking; [WARN]s are notes,
                           e.g. daily tables that don't add up to the weekly
                           table. The files are produced either way.)

Usage:
    python run_monthly_reports.py \
        --source MTN_F5_ATTACKS_SOURCE.xlsx \
        --pptx-template MTN_Security_Metrics_Report_SOURCE.pptx \
        --xlsx-template MTN_F5_WAF_Weekly_Report_Temp.xlsx \
        --out-dir ./monthly_output \
        [--month 2026-09]            # default: the newest month the source fully covers
        [--certs Certificate.xlsx]   # certificate slides are skipped if it's missing
        [--activity-dir activity_reports]   # weekly activity reports -> monthly activity report

Output files are named after the month, e.g.:
    monthly_output/MTN_Security_Metrics_Report_September2026.pptx
    monthly_output/MTN_F5_WAF_Monthly_Report_September2026.xlsx

generate_monthly_reports() is what the command line, the .bat file and the
web app all call - one code path, whichever way it's run.
"""

import argparse
import gc
import os
from dataclasses import dataclass, field
from datetime import date, timedelta

import openpyxl
from pptx import Presentation

import generate_weekly_report as wk
import generate_monthly_report as mon
import generate_monthly_excel as mx
import generate_activity_report as act

KPI_JUMP_WARN_RATIO = 0.5      # warn if a site's month moved more than +-50% vs last month


@dataclass
class MonthlyResult:
    month: tuple              # (year, month)
    month_label: str          # "September 2026"
    pptx_path: str
    xlsx_path: str
    ok: bool                  # False if any [FAIL]
    log: list = field(default_factory=list)       # every line printed
    warnings: list = field(default_factory=list)  # the [WARN] lines
    checks: list = field(default_factory=list)    # the [PASS]/[WARN]/[FAIL] lines
    activity_path: str = None                     # the monthly activity report, if one was built


def _days_text(days):
    """[28 Sep, 29 Sep, 30 Sep] -> '28-30 Sep'."""
    if not days:
        return ""
    a, b = days[0], days[-1]
    if a == b:
        return f"{a.day} {a:%b}"
    return f"{a.day}-{b.day} {b:%b}" if a.month == b.month else f"{a.day} {a:%b}-{b.day} {b:%b}"


def month_options(sheetnames, today=None):
    """
    For the web page's dropdown, from sheet NAMES only: every month the
    source has sheets for, newest first, with a note when some of its days
    aren't in the file yet. The default is the newest month it fully covers.
    """
    class _Names:
        pass
    _Names.sheetnames = list(sheetnames)
    weeks = wk.group_sheets_by_week(_Names)
    covs = [c for c in mon.available_months(weeks, today) if c.any_data]
    if not covs:
        return None
    default = mon.pick_target_month(covs)
    options = []
    for c in reversed(covs):
        label = mon.month_label(c.ym)
        if not c.complete:
            gaps = sorted(set(d for v in c.missing.values() for d in v))
            have = [d for d in mon.month_days(c.ym) if d not in gaps]
            if len(gaps) <= 10:
                label += f"  (no sheets yet for {_days_text(gaps)})"
            else:
                label += f"  (only {_days_text(have)} in the file)" if have else "  (incomplete)"
        options.append({"value": f"{c.ym[0]}-{c.ym[1]:02d}", "label": label, "selected": c.ym == default.ym})
    return {"options": options, "default": default.ym}


def _load_certs(certs_path, ym, out):
    if not certs_path or not os.path.exists(certs_path):
        out(f"  [WARN] {certs_path or 'Certificate.xlsx'} not found - certificate slides skipped")
        return None
    certs = wk.read_certificates(certs_path, ym[0])
    first = mon.month_days(ym)[0]
    for site in wk.SITES:
        info = certs.get(site)
        if info is None:
            out(f"  [WARN] {site}: device {wk.CERT_NODES[site]} not found in {os.path.basename(certs_path)}"
                f" - no certificate slides for {site}")
            continue
        out(f"  {site}: {len(info['ssl'])} SSL + {len(info['device'])} device certificate(s) "
            f"expiring in {ym[0]} ({info['node']})")
        for c in info["unreadable"]:
            out(f"  [WARN] {site}: couldn't read the expiry date '{c.expiry_text}' for {c.name}"
                f" - left off the slide, check Certificate.xlsx")
        checked = info.get("checked")
        if isinstance(checked, date) and checked < first - timedelta(days=7):
            out(f"  [WARN] {site}: certificate check is dated {checked:%d %b %Y}, before this month - "
                f"is Certificate.xlsx up to date?")
    return certs


def generate_monthly_reports(source, pptx_template, xlsx_template, out_dir,
                             month=None, certs="Certificate.xlsx", log=print,
                             activity_files=None, activity_names=None):
    """
    Builds the monthly deck + Excel for `month` ((year, month) or None = the
    newest month the source fully covers), checks them, and returns a
    MonthlyResult. `log` gets every progress line. activity_files: the
    month's weekly activity report workbooks (any order, extra weeks are
    ignored) -> also builds the monthly activity report; activity_names:
    their original file names, for messages (optional).
    """
    lines = []

    def out(msg=""):
        lines.append(msg)
        log(msg)

    seen = set()

    def warn(text):
        if text not in seen:          # a split week is checked for both months it touches
            seen.add(text)
            out(f"  [WARN] {text}")

    os.makedirs(out_dir, exist_ok=True)
    ctx = mon.load_month_context(source, month)
    ym = ctx.target
    label = mon.month_label(ym)
    tag = f"{mon.month_name(ym)}{ym[0]}"
    pptx_out = os.path.join(out_dir, f"MTN_Security_Metrics_Report_{tag}.pptx")
    xlsx_out = os.path.join(out_dir, f"MTN_F5_WAF_Monthly_Report_{tag}.xlsx")

    out("=" * 60)
    out(f"MONTHLY REPORT: {label}")
    out("=" * 60)
    for site, days in ctx.coverage.missing.items():
        if days:
            warn(f"{site}: no sheet yet for {_days_text(days)} - those days aren't counted, "
                 f"so {label} is incomplete")
    out(f"Trend chart months: {', '.join(mon.month_label(t) for t in ctx.trend)}")
    for t, why in ctx.skipped:
        out(f"  (left off the trend chart: {mon.month_label(t)} - {why})")

    out()
    out("Month totals (whole weeks from the Weekly tables, split weeks from the day tables):")
    site_trends = {}
    for site in wk.SITES:
        history = [mon.compute_site_month(ctx, site, t, warn=warn) for t in ctx.trend[:-1]]
        cur = mon.compute_site_month(ctx, site, ym, detail=True, warn=warn)
        site_trends[site] = history + [cur]
        out(f"  {site}: Requests Inspected={cur.requests_inspected:,}  WAF Blocks={cur.waf_blocks:,}  "
            f"Pool={cur.pool_pct:.2f}% ({cur.pool_available}/{cur.pool_total}"
            f"{', ' + cur.pool_date.strftime('%d %b') if cur.pool_date else ''})")
        for p in cur.parts:
            what = "no sheet" if p.missing else f"{p.requests:,} / {p.waf_blocks:,} ({p.source} tables)"
            out(f"      {p.label}: {what}")
    total_trend = [mon.combine(t, [site_trends[s][i] for s in wk.SITES])
                   for i, t in enumerate(ctx.trend)]
    total = total_trend[-1]
    out(f"  TOTAL: Requests Inspected={total.requests_inspected:,}  WAF Blocks={total.waf_blocks:,}  "
        f"Pool={total.pool_pct:.2f}%")

    out()
    out("Node changes / certificates:")
    node_changes = {}
    for site in wk.SITES:
        changes, ambiguous = mon.month_node_changes(ctx, site, ym)
        node_changes[site] = changes
        out(f"  {site}: {len(changes)} node change(s)")
        if ambiguous:
            days = sorted({d for d, _ in ambiguous})
            warn(f"{site}: {len(ambiguous)} node(s) listed under more than one status on the same day "
                 f"({', '.join(d.strftime('%d %b') for d in days[:6])}{' ...' if len(days) > 6 else ''})"
                 f" - left off the slides, check them by hand")
    cert_data = _load_certs(certs, ym, out)

    mon.fill_monthly_pptx(pptx_template, pptx_out, ym, total_trend, site_trends,
                          node_changes=node_changes, certs=cert_data)
    out(f"Wrote {pptx_out}")
    mx.write_monthly_excel(xlsx_template, xlsx_out, ym, site_trends)
    out(f"Wrote {xlsx_out}")

    out()
    out("=" * 60)
    out("CHECKS (READ THIS PART)")
    out("=" * 60)
    checks = run_monthly_checks(ctx, site_trends, total_trend, pptx_out, xlsx_out)
    for c in checks:
        out(c)
    ok = not any(c.startswith("[FAIL]") for c in checks)

    activity_path = None
    if activity_files:
        # the WAF month data isn't needed any more - free it before opening the weekly
        # activity workbooks, so the two never sit in memory together (Render's 512 MB)
        del ctx, site_trends, total_trend
        gc.collect()
        out()
        out("=" * 60)
        out(f"MONTHLY ACTIVITY REPORT ({len(activity_files)} weekly file(s) given)")
        out("=" * 60)
        res = act.build_activity_report(activity_files, ym, os.path.join(out_dir, act.activity_file_name(ym)),
                                        warn=warn, names=activity_names)
        if res is not None:
            activity_path = res.path
            for key, name in res.weeks_used:
                out(f"  used {key[0]:%d %b} - {key[1]:%d %b}: {name}")
            out(f"  health check sheets: " + ", ".join(
                f"Week {i} ({act._span(days)})" for i, days in enumerate(res.health_weeks, start=1)))
            out(f"  {res.cases} case(s) in the Overview, {res.case_sheets} case screenshot sheet(s); "
                f"devices: DNS {len(res.stats['DNS'].nodes)}, LTM {len(res.stats['LTM'].nodes)}")
            out(f"Wrote {activity_path}")

    return MonthlyResult(
        month=ym, month_label=label, pptx_path=pptx_out, xlsx_path=xlsx_out, ok=ok, log=lines,
        warnings=[l.strip() for l in lines if "[WARN]" in l and not l.startswith("[WARN]")],
        checks=checks, activity_path=activity_path,
    )


def run_monthly_checks(ctx, site_trends, total_trend, pptx_path, xlsx_path):
    """Re-opens both files and checks them against the month's numbers."""
    res = []

    def check(label, ok, detail=""):
        res.append(f"[{'PASS' if ok else 'FAIL'}] {label}" + (f" - {detail}" if detail and not ok else ""))

    ym = ctx.target
    for site in wk.SITES:
        cur = site_trends[site][-1]
        present = [p for p in cur.parts if not p.missing]
        check(f"{site}: has sheets in {mon.month_label(ym)}", bool(present),
              "no sheet for this site in any week of the month")
        if not present:
            continue
        check(f"{site}: Requests Inspected is not 0", cur.requests_inspected > 0,
              "no Request Throughput table was read - check the table titles")
        check(f"{site}: WAF Blocks is not 0", cur.waf_blocks > 0,
              "no WAF Attack Breakdown table was read - check the table titles")
        check(f"{site}: pool member status found", cur.pool_total > 0,
              "no Pool Member Status table was read for any day of the month")
        prev = site_trends[site][-2] if len(site_trends[site]) > 1 else None
        if prev and prev.ym == mon.prev_month(ym):
            for name, a, b in (("Requests Inspected", cur.requests_inspected, prev.requests_inspected),
                               ("WAF Blocks", cur.waf_blocks, prev.waf_blocks)):
                if b and abs(a - b) / b > KPI_JUMP_WARN_RATIO:
                    res.append(f"[WARN] {site}: {name} moved {(a - b) / b * 100:+.0f}% vs "
                               f"{mon.month_name(prev.ym)} ({b:,} -> {a:,}) - worth a second look")

    # Excel totals = PowerPoint WAF Blocks
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    try:
        for site in wk.SITES:
            ws = wb[mx.wx.SHEET_MAP[site][0]]
            xl_total = None
            for row in ws.iter_rows(values_only=True):
                if len(row) >= 9 and isinstance(row[6], str) and row[6].strip().upper() == "TOTAL":
                    xl_total = row[8]
            want = site_trends[site][-1].waf_blocks
            check(f"{site}: Excel attack-type total matches the deck's WAF Blocks ({want:,})",
                  xl_total == want, f"Excel says {xl_total}")
    finally:
        wb.close()

    # the TOTAL dashboard shows the combined figures
    prs = Presentation(pptx_path)
    texts = {sh.text_frame.text.strip() for sl in list(prs.slides)[1:2] for sh in sl.shapes if sh.has_text_frame}
    total = total_trend[-1]
    check("Deck: TOTAL dashboard shows the combined Requests Inspected",
          f"{total.requests_inspected:,}" in texts, "value not found on slide 2")
    check("Deck: TOTAL dashboard shows the combined WAF Blocks",
          f"{total.waf_blocks:,}" in texts, "value not found on slide 2")
    return res


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", required=True, help="MTN_F5_ATTACKS_SOURCE.xlsx")
    ap.add_argument("--pptx-template", required=True, help="MTN_Security_Metrics_Report_SOURCE.pptx")
    ap.add_argument("--xlsx-template", required=True, help="MTN_F5_WAF_Weekly_Report_Temp.xlsx")
    ap.add_argument("--out-dir", default=".", help="Folder to write outputs into (created if missing)")
    ap.add_argument("--month", default=None,
                    help="e.g. 2026-09 or 'Sep 2026' - defaults to the newest month the source fully covers")
    ap.add_argument("--certs", default="Certificate.xlsx",
                    help="Certificate workbook for the certificate slides (skipped with a warning if missing)")
    ap.add_argument("--activity-dir", default=None,
                    help="Folder holding the month's weekly activity reports (.xlsx) - builds the monthly "
                         "activity report too (skipped if the folder is missing or empty)")
    args = ap.parse_args()
    month = mon.parse_month(args.month) if args.month else None
    activity = None
    if args.activity_dir:
        if os.path.isdir(args.activity_dir):
            activity = sorted(os.path.join(args.activity_dir, f) for f in os.listdir(args.activity_dir)
                              if f.lower().endswith(".xlsx") and not f.startswith("~$"))
        if not activity:
            print(f"(No weekly activity reports in '{args.activity_dir}' - the monthly activity report is skipped.)")
    result = generate_monthly_reports(args.source, args.pptx_template, args.xlsx_template, args.out_dir,
                                      month=month, certs=args.certs, activity_files=activity)
    print()
    if result.ok:
        print("All checks passed." + (" Glance at the [WARN] notes above." if result.warnings else ""))
    else:
        print("At least one check FAILED - look at the [FAIL] lines before sending anything.")
    raise SystemExit(0 if result.ok else 1)


if __name__ == "__main__":
    main()
